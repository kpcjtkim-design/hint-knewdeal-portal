#!/usr/bin/env python3
"""2기 일정표 엑셀(한 판 또는 여러 판)을 합쳐 timetable-seed.json·timetable-reference.json을 만듭니다.

사용: python -I tools/build-timetable-seed.py --sheet 1014 <일정표.xlsx> <원본 파일명> [--sheet 1007 <이전 일정표.xlsx> <원본 파일명>]
      --version YYYYMMDD [--conflicts 충돌.csv] [--check] [--no-align] [--force] [--out 폴더(시험용, 저장소 파일 그대로)]
2기 첫 기준(20261014) 만든 명령: --sheet 1014 <261014 파일> "261014_K-뉴딜_아카데미_2기 일정표_직무기초.xlsx"
      --sheet 1007 <261007 파일> "261007_K-뉴딜_아카데미_2기 일정표.xlsx" --version 20261014 --conflicts tt2-conflicts.csv

- 반·기간·휴일은 cohort.json(2기 16반, 2026-10-28~2027-01-22)에서 읽습니다. 엑셀 반 열은 2행 장소·4행 과정(트랙 번호)으로 찾고,
  맞는 반이 없거나 둘 이상이면 멈춥니다. 대구상공회의소 열은 사용자 결정으로 2기에서 뺐습니다.
- 머리글: A열 '날짜' 행 아래 '과정명·시간·강사명' 행, 그 다음 행부터 날짜별 한 줄(주말·휴일 줄 포함). 반 묶음은 '시간'·'강사명'이 이어지는 3칸.
- 여러 판 합치기(사용자 결정 2026-10-19 '두 개 절충'): 반·날짜 칸마다 과목·시간·강사를 따로 봅니다. 한쪽만 있으면 그 값,
  둘 다 같으면 그 값(출처 both), 둘 다 있고 다르면 첫 --sheet(최신) 값 + 충돌 표시(entry.review, --conflicts CSV). 출처는 entry.origin.
  둘 다 과목이 비었는데 시간이 있으면 '직무특화 (세부 과목 미정)'으로 둡니다(과목을 지어내지 않음).
- 확인된 엑셀 오류는 ALIGN에서 바로잡습니다(--no-align이면 그대로 읽어 주말·휴일 칸이 있으면 멈춤).
- 1기 엑셀(260818 전체 일정표, 17반) 모드는 2기 브랜치에서 뺐습니다(cohort1 브랜치·git 기록). tools/timetable-overrides.json(1기 골든벨)은 읽지 않습니다.
"""
import json, re, sys, csv, datetime, pathlib, collections
import openpyxl
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
ROOT = pathlib.Path(__file__).resolve().parent.parent
SEED, REFERENCE, HISTORY = ROOT / 'timetable-seed.json', ROOT / 'timetable-reference.json', ROOT / 'timetable-seeds'
COHORT = json.loads((ROOT / 'cohort.json').read_text(encoding='utf-8'))
CLASSES = {c['id']: c for c in COHORT['classes']}
HOLIDAYS = {h['date']: h['name'] for h in COHORT['holidays']}
# 반 → (2행 장소 정규식, 4행 트랙 번호). 과정(모빌리티 AI·제조지능화)은 cohort.json 반 과정과 4행이 같아야 합니다. 같은 장소·과정이 둘이면 트랙 번호로 가릅니다.
CLASS_PLACES = {'1': (r'서울대', None), '2': (r'한양대', None), '3': (r'후인원', None), '4': (r'후인원', None), '5': (r'청주', None),
                '6': (r'생산성본부|대구무역회관', None), '7': (r'마이크로웨이브', None), '8': (r'마이크로웨이브', None), '9': (r'부산대', None),
                '10': (r'아르피나', None), '11': (r'아르피나', '1'), '12': (r'아르피나', '2'), '13': (r'울산', None), '14': (r'전남대', None),
                '15': (r'광주공장|기아오토랜드', '1'), '16': (r'광주공장|기아오토랜드', '2')}
EXCLUDED = r'대구상공회의소'  # 사용자 결정: 2기 16반에서 뺀 반
# 확인된 엑셀 오류 바로잡기(사용자 확인 '맞춰' 2026-10-19). 261007판 15반(기아오토랜드A) 과정명 칸만 2026-12-02부터 한 줄씩 아래로 밀림:
# 금요일 과목이 토요일 줄, 월요일 칸 빔, 12/25·12/31 휴일 줄에 과목, 수료식이 날짜 없는 줄. 시간·강사명 칸은 제자리라 과정명만 한 줄 위로 당깁니다.
ALIGN = [{'sheet': '1007', 'classId': '15', 'field': 'subject', 'from': '2026-12-02', 'up': 1,
          'why': '261007판 15반 과정명 칸이 12/02부터 한 줄 밀림(토요일·휴일 줄에 과목, 수료식이 날짜 없는 줄)'}]
# 같은 과목의 다른 표기(띄어쓰기·옛 이름) → 한 이름. 비교는 공백을 뺀 글자로 합니다.
ALIAS = {'입교식-프로그램': '입교식', 'C프로그래밍': 'C 프로그래밍', 'MCU프로그래밍': 'MCU 프로그래밍', '모빌리티SW개론및프로세스': '모빌리티 SW 개론 및 프로세스',
         '자동차통신시스템': '자동차 통신 시스템', '딥러닝기반영상인식': '딥러닝 기반 영상인식', '직무특화컨퍼런스1': '직무특화 컨퍼런스 1', '직무특화컨퍼런스2': '직무특화 컨퍼런스 2',
         'AI기반스마트팩토리설계및운영실습': 'AI기반 스마트팩토리 설계 및 운영 실습', 'AI기반스마트팩토리설계및운영기술': 'AI기반 스마트팩토리 설계 및 운영 기술'}
UNDECIDED = '직무특화 (세부 과목 미정)'
MODULES = ['공통행사', '디지털스킬', '소프트스킬', '직무특화', '공장견학', '분해조립', '진로취업', '성장지원', '현직자특강']
WEAK = re.compile(r'^(미정|-|tbd)$', re.I)  # 강사 '미정'은 다른 판의 이름이 있으면 그 이름을 씁니다(충돌로 치지 않음)
FIELDS = {'subject': '과목', 'hours': '시간', 'instructor': '강사'}


def compact(text):
    return re.sub(r'\s+', '', str(text or ''))


def canon(text):
    t = re.sub(r'\s+', ' ', str(text or '')).strip()
    return ALIAS.get(compact(t), t)


def classify(s):
    """정리한 과목명 → module·title·event(입교식·수료식·공장견학·분해조립)·venue·online·kind."""
    out = {'module': '직무특화', 'title': s, 'event': '', 'venue': '', 'online': False, 'kind': 'class'}
    k = compact(s)
    if k == '휴일':
        return {**out, 'module': '휴일', 'kind': 'holiday'}
    if m := re.fullmatch(r'(\S+공장)\(견학\+이론\)', k):
        return {**out, 'module': '공장견학', 'title': f'{m[1]} · 견학·이론', 'event': '공장견학', 'venue': m[1]}
    if k == '공장견학(이론)':  # 아르피나B·C: 견학 없이 이론만(엑셀 '울산공장(견학)' 0일)
        return {**out, 'module': '공장견학', 'title': '공장견학 · 이론', 'event': '공장견학(이론)'}
    if m := re.fullmatch(r'(.+)-분해조립', k):
        return {**out, 'module': '분해조립', 'title': f'{m[1]} · 분해조립', 'event': '분해조립', 'venue': m[1]}
    if s in ('입교식', '수료식'):
        return {**out, 'module': '공통행사', 'event': s}
    if s == '프로그램':
        return {**out, 'module': '공통행사'}
    if s in ('디지털스킬', '소프트스킬', '진로취업', '성장지원'):
        return {**out, 'module': s}
    if k == '현직자특강(비대면)':
        return {**out, 'module': '현직자특강', 'title': '현직자특강', 'online': True}
    if m := re.fullmatch(r'직무특화\(비대면/(.+)\)', s):
        return {**out, 'title': m[1].strip(), 'online': True}
    return out


def course_of(track):
    t = str(track or '')
    return '모빌리티 AI' if re.search(r'모[빌비]리티', t) else '제조지능화' if t.strip().startswith('제조지능화') else ''


def read_sheet(tag, path):
    """엑셀 한 판 → {'groups': {반: 열 글자}, 'cells': {반: [{row,date,subject,hours,instructor}]}, 'raw': 원본 칸(장소 등)}."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    if str(ws.cell(2, 1).value or '').strip() != '교육장소':
        sys.exit(f'[{tag}] 2행 A열이 교육장소가 아닙니다. 일정표 모양을 확인해 주세요.')
    head = next((r for r in range(1, 60) if str(ws.cell(r, 1).value or '').strip() == '날짜'), None)
    label = next((r for r in range(head or 1, (head or 1) + 4) if any(ws.cell(r, c).value == '강사명' for c in range(1, ws.max_column + 1))), None) if head else None
    if not label:
        sys.exit(f'[{tag}] 날짜·과정명·시간·강사명 머리글을 찾지 못했습니다.')
    groups, excluded = {}, []
    for c in range(1, ws.max_column - 1):
        if ws.cell(label, c + 1).value != '시간' or ws.cell(label, c + 2).value != '강사명':
            continue
        place, track = str(ws.cell(2, c).value or '').strip(), str(ws.cell(4, c).value or '').strip()
        if re.search(EXCLUDED, place):
            excluded.append(f'{get_column_letter(c)}({place})')
            continue
        course, no = course_of(track), (re.findall(r'(\d+)\s*$', track) or [None])[0]
        hits = [cid for cid, (rx, n) in CLASS_PLACES.items() if re.search(rx, place) and CLASSES[cid]['course'] == course and (n is None or n == no)]
        if len(hits) != 1:
            sys.exit(f'[{tag}] {get_column_letter(c)}열({place} · {track})에 맞는 2기 반이 {len(hits)}개입니다: {hits}. CLASS_PLACES를 확인해 주세요.')
        if hits[0] in groups:
            sys.exit(f'[{tag}] {hits[0]}반 열이 둘입니다({groups[hits[0]]}, {get_column_letter(c)}).')
        groups[hits[0]] = c
    if set(groups) != set(CLASSES):
        sys.exit(f'[{tag}] 반 열이 모자랍니다: {sorted(set(CLASSES) - set(groups), key=int)}')
    cells = {cid: [] for cid in groups}
    for r in range(label + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        date = v.date().isoformat() if isinstance(v, datetime.datetime) else None
        if v not in (None, '') and not date:
            sys.exit(f'[{tag}] {r}행 날짜를 읽을 수 없습니다: {v!r}')
        for cid, c in groups.items():
            s, h, n = (ws.cell(r, c + i).value for i in range(3))
            h = h if isinstance(h, (int, float)) and not isinstance(h, bool) else float(h) if re.fullmatch(r'\d+(\.\d+)?', str(h or '').strip()) else None
            cells[cid].append({'row': r, 'date': date, 'subject': canon(s), 'raw': str(s or '').strip(), 'hours': h, 'instructor': re.sub(r'\s+', ' ', str(n or '')).strip()})
    totals = {}  # 위쪽 합계표(261007판 6~28행): 항목 → {반: 일수}
    for r in range(3, head):
        name = str(ws.cell(r, 1).value or '').strip()
        if name and name not in ('구분', '교육장소', '합계'):
            totals[name] = {cid: ws.cell(r, c).value for cid, c in groups.items() if isinstance(ws.cell(r, c).value, (int, float))}
    places = {cid: str(ws.cell(2, c).value or '').strip() for cid, c in groups.items()}
    return {'tag': tag, 'groups': {cid: get_column_letter(c) for cid, c in groups.items()}, 'cells': cells, 'excluded': excluded, 'totals': totals, 'places': places}


def align(sheet, report):
    for rule in ALIGN:
        if rule['sheet'] != sheet['tag']:
            continue
        rows, f = sheet['cells'][rule['classId']], rule['field']
        start = next(i for i, x in enumerate(rows) if x['date'] and x['date'] >= rule['from'])
        moved = []
        for i in range(start, len(rows)):
            nxt = rows[i + rule['up']] if i + rule['up'] < len(rows) else {f: '', 'raw': ''}
            if (rows[i][f] or nxt[f]) and rows[i][f] != nxt[f]:
                moved.append(f"{rows[i]['date'] or str(rows[i]['row']) + '행'}: {rows[i][f] or '(빈칸)'} → {nxt[f] or '(빈칸)'}")
            rows[i] = {**rows[i], f: nxt[f], 'raw': nxt['raw']}
        report.append({**{k: rule[k] for k in ('sheet', 'classId', 'field', 'from', 'up', 'why')}, 'changed': len(moved), 'moves': moved})


def offday_problems(sheet):
    out = []
    for cid, rows in sheet['cells'].items():
        for x in rows:
            if not (x['subject'] or x['hours']):
                continue
            d = x['date']
            weekend = d and datetime.date.fromisoformat(d).weekday() >= 5
            if not d or weekend or d in HOLIDAYS or not (COHORT['period']['start'] <= d <= COHORT['period']['end']):
                if d in HOLIDAYS and compact(x['subject']) == '휴일':
                    continue
                out.append(f"[{sheet['tag']}] {cid}반 {d or str(x['row']) + '행(날짜 없음)'}: {x['subject'] or '(과목 없음)'} {x['hours'] or ''}")
    return out


def training_days():
    d, end, out = datetime.date.fromisoformat(COHORT['period']['start']), datetime.date.fromisoformat(COHORT['period']['end']), []
    while d <= end:
        if d.weekday() < 5 and d.isoformat() not in HOLIDAYS:
            out.append(d.isoformat())
        d += datetime.timedelta(days=1)
    return out


def merge(field, values):
    """values: [(판, 값)] 우선순위 순. → (값, 출처, 충돌 여부). 값이 같은지는 공백을 뺀 글자(시간은 숫자)로 봅니다."""
    key = (lambda v: float(v)) if field == 'hours' else compact
    present = [(t, v) for t, v in values if v not in (None, '')]
    if field == 'instructor' and any(not WEAK.match(v) for _, v in present):
        present = [(t, v) for t, v in present if not WEAK.match(v)]
    if not present:
        return ('' if field != 'hours' else None), '', False
    tag, value = present[0]
    same = [t for t, v in present if key(v) == key(value)]
    return value, ('both' if len(same) > 1 else tag), len(same) < len(present)


def args_of(argv):
    sheets, opts, i = [], {}, 0
    while i < len(argv):
        a = argv[i]
        if a == '--sheet':
            sheets.append(tuple(argv[i + 1:i + 4])); i += 4
        elif a in ('--version', '--conflicts', '--out'):
            opts[a[2:]] = argv[i + 1]; i += 2
        elif a in ('--check', '--no-align', '--force'):
            opts[a[2:]] = True; i += 1
        else:
            sys.exit(f'알 수 없는 인자: {a}\n{__doc__}')
    if not sheets or any(len(s) != 3 for s in sheets) or 'version' not in opts:
        sys.exit(__doc__)
    return sheets, opts


def main():
    specs, opts = args_of(sys.argv[1:])
    version, primary = opts['version'], specs[0][0]
    old = json.loads(SEED.read_text(encoding='utf-8')) if SEED.exists() else {}
    same_cohort = old.get('cohort') == COHORT['id']
    sheets, aligned = [read_sheet(tag, path) for tag, path, _ in specs], []
    for s in sheets:
        print(f"[{s['tag']}] 반 열: " + ' · '.join(f"{cid} {s['groups'][cid]}" for cid in CLASSES) + (f" · 뺀 열 {', '.join(s['excluded'])}" if s['excluded'] else ''))
        if not opts.get('no-align'):
            align(s, aligned)
    for a in aligned:
        print(f"[{a['sheet']}] {a['classId']}반 {FIELDS[a['field']]} 칸 {a['from']}부터 {a['up']}줄 당김 · 바뀐 칸 {a['changed']}개 ({a['why']})")
        print('  ' + '\n  '.join(a['moves']))
    problems = [p for s in sheets for p in offday_problems(s)]
    if problems:
        sys.exit('주말·휴일·기간 밖·날짜 없는 줄에 수업이 있습니다. ALIGN 규칙이나 엑셀을 먼저 확인해 주세요:\n  ' + '\n  '.join(problems))

    days, tags = training_days(), [s['tag'] for s in sheets]
    lectures = list(old.get('lectures', [])) if same_cohort else []
    by_title = {(l['course'], compact(l['title'])): l for l in lectures}
    numbers = [int(m[1]) for l in lectures if (m := re.fullmatch(r't2-(\d+)', l['id']))]
    classes, conflicts, summary = {}, [], {}
    for cid, meta in CLASSES.items():
        course, entries, counter = meta['course'], [], collections.Counter()
        stat = collections.Counter()
        cells = {s['tag']: {x['date']: x for x in s['cells'][cid] if x['date']} for s in sheets}
        for date in sorted({d for t in tags for d, x in cells[t].items() if x['subject'] or x['hours']} | set(days)):
            if date in HOLIDAYS:
                continue
            got = {t: cells[t].get(date, {}) for t in tags}
            subject, ssrc, sc = merge('subject', [(t, got[t].get('subject', '')) for t in tags])
            hours, hsrc, hc = merge('hours', [(t, got[t].get('hours')) for t in tags])
            name, nsrc, nc = merge('instructor', [(t, got[t].get('instructor', '')) for t in tags])
            if not subject and not hours:
                stat['missing'] += 1
                print(f'  ! {cid}반 {date}: 어느 판에도 수업 칸이 없습니다.')
                continue
            for field, flag, chosen in (('subject', sc, subject), ('hours', hc, hours), ('instructor', nc, name)):
                if flag:
                    conflicts.append({'date': date, 'classId': cid, 'field': FIELDS[field], **{t: got[t].get(field, '') for t in tags}, 'chosen': chosen})
            info = classify(subject) if subject else {'module': '직무특화', 'title': UNDECIDED, 'event': '', 'venue': '', 'online': False, 'kind': 'class'}
            stat['subject:' + (ssrc or 'none')] += 1
            if info['kind'] == 'holiday':
                continue
            lecture = by_title.get((course, compact(info['title'])))
            if not lecture:  # 새 강의는 임시 열쇠로 두었다가 아래에서 과정·첫 날짜 순으로 t2-번호를 붙입니다(이전 2기 시드 번호는 유지)
                lecture = {'id': f"new|{course}|{compact(info['title'])}", 'course': course, 'module': info['module'], 'title': info['title'], 'days': 0}
                lectures.append(lecture); by_title[(course, compact(info['title']))] = lecture
            counter[lecture['id']] += 1
            src = primary if ssrc in ('both', '') or not got[ssrc] else ssrc  # sourceRow·sourceText는 과목을 가져온 판 기준
            src = src if got[src] else next(t for t in tags if got[t])
            review = [f for f, flag in (('subject', sc), ('hours', hc), ('instructor', nc)) if flag]
            entries.append({'id': f'x2-{cid}-{date}', 'date': date, 'course': course, 'module': info['module'], 'lectureId': lecture['id'],
                            'title': info['title'], 'day': counter[lecture['id']], 'hours': (int(hours) if float(hours).is_integer() else hours) if hours else 0,
                            'start': '', 'end': '', 'instructorId': '', 'instructorName': name, 'venue': info['venue'], 'online': info['online'], 'note': '',
                            'kind': 'class', 'event': info['event'], 'origin': {'subject': ssrc, 'hours': hsrc, 'instructor': nsrc},
                            **({'review': review} if review else {}), 'sourceRow': got[src]['row'], 'sourceText': got[src].get('raw', '')})
        for date, name in HOLIDAYS.items():  # cohort.json 휴일(평일)은 모든 반에 휴일 칸으로 둡니다(대체휴일 순연·이동이 피해 가도록)
            if datetime.date.fromisoformat(date).weekday() < 5 and COHORT['period']['start'] <= date <= COHORT['period']['end']:
                entries.append({'id': f'holiday-{cid}-{date}', 'date': date, 'course': course, 'module': '휴일', 'lectureId': '', 'title': name, 'day': 0, 'hours': 0,
                                'start': '', 'end': '', 'instructorId': '', 'instructorName': '', 'venue': '', 'online': False, 'note': '', 'kind': 'holiday', 'event': ''})
        entries.sort(key=lambda e: (e['date'], e['id']))
        classes[cid] = {'classId': cid, 'revision': 0, 'entries': entries}
        lessons = [e for e in entries if e['kind'] != 'holiday']
        covered = {e['date'] for e in lessons}
        summary[cid] = {'days': len(covered & set(days)), 'offDays': sorted(covered - set(days)), 'missing': sorted(set(days) - covered),
                        'hours': sum(e['hours'] for e in lessons), 'undecided': sum(e['title'] == UNDECIDED for e in lessons),
                        'subject': {k.split(':')[1]: v for k, v in stat.items() if k.startswith('subject:')},
                        'conflicts': collections.Counter(c['field'] for c in conflicts if c['classId'] == cid)}
    for l in lectures:
        l['days'] = max([sum(e['lectureId'] == l['id'] for e in c['entries']) for c in classes.values()] + [0]) or l['days']
    lectures = [l for l in lectures if l['days']]
    first = {}
    for c in classes.values():
        for e in c['entries']:
            first.setdefault(e['lectureId'], e['date'])
    lectures.sort(key=lambda l: (COHORT['courses'].index(l['course']), first.get(l['id'], '9999'), MODULES.index(l['module']), l['title']))
    rename = {}
    for l in lectures:
        if l['id'].startswith('new|'):
            numbers.append(max(numbers, default=0) + 1); rename[l['id']] = l['id'] = f't2-{numbers[-1]}'
    for c in classes.values():
        for e in c['entries']:
            e['lectureId'] = rename.get(e['lectureId'], e['lectureId'])

    print('\n반 · 교육일/60 · 시간 · 과목 출처 · 미정 · 충돌')
    for cid, s in summary.items():
        print(f"  {cid}반 · {s['days']}/{len(days)}일 · {s['hours']}시간 · " + ', '.join(f'{k} {v}' for k, v in sorted(s['subject'].items()))
              + f" · 미정 {s['undecided']} · 충돌 " + (', '.join(f'{k} {v}' for k, v in s['conflicts'].items()) or '0')
              + (f" · 빠진 교육일 {s['missing']}" if s['missing'] else '') + (f" · 교육일 밖 {s['offDays']}" if s['offDays'] else ''))

    # 공장견학·분해조립 장소·날짜(일정에서), 교육장소(각 판 2행), 합계표(있으면)와 대조
    facilities, checks = [], []
    for cid in CLASSES:
        rows = [e for e in classes[cid]['entries'] if e['event'] in ('공장견학', '공장견학(이론)', '분해조립')]
        visit = [e for e in rows if e['event'] == '공장견학'] or [e for e in rows if e['event'] == '공장견학(이론)']
        dis = [e for e in rows if e['event'] == '분해조립']
        places = {s['tag']: s['places'][cid] for s in sheets}
        facilities.append({'classId': cid, 'factory': visit[0]['venue'] if visit and visit[0]['venue'] else '공장견학 이론만(견학 일정 없음)' if visit else '',
                           'factoryDates': [e['date'] for e in visit], 'disassembly': dis[0]['venue'] if dis else '', 'disassemblyDates': [e['date'] for e in dis],
                           'venue': f"{CLASSES[cid]['area']} {CLASSES[cid]['venue']}", 'address': max(places.values(), key=len), 'places': places})
        for s in sheets:
            for label, counts in s['totals'].items():
                if not re.search(r'공장|분해조립|견학', label) or cid not in counts:
                    continue
                k = compact(canon(label))
                got = sum(1 for x in s['cells'][cid] if x['date'] and compact(x['subject']) == k)
                if abs(got - counts[cid]) > 0.01 and not (got == 1 and counts[cid] == 0.5):
                    checks.append(f"[{s['tag']}] {cid}반 합계표 '{label}' {counts[cid]}일 · 일정 칸 {got}일")
    if checks:
        print('\n합계표와 일정 칸 수가 다른 곳(바로잡기 뒤 기준):\n  ' + '\n  '.join(checks))
    print('\n공장견학 · 분해조립')
    for f in facilities:
        print(f"  {f['classId']}반 · {f['factory'] or '-'} {','.join(f['factoryDates'])} · {f['disassembly'] or '-'} {','.join(f['disassemblyDates'])}")

    if opts.get('conflicts'):
        with open(opts['conflicts'], 'w', encoding='utf-8-sig', newline='') as fh:
            w = csv.writer(fh)
            w.writerow(['날짜', '반', '항목'] + [f'{t}판 값' for t in tags] + ['선택(기본: 첫 판)'])
            for c in sorted(conflicts, key=lambda c: (int(c['classId']), c['date'], c['field'])):
                w.writerow([c['date'], c['classId'], c['field']] + [c[t] if c[t] not in (None,) else '' for t in tags] + [c['chosen']])
        print(f"\n충돌 {len(conflicts)}건 → {opts['conflicts']}")
    else:
        print(f'\n충돌 {len(conflicts)}건(--conflicts로 CSV 저장)')

    names = {tag: name for tag, _, name in specs}
    dates = [e['date'] for c in classes.values() for e in c['entries'] if e['kind'] != 'holiday']
    seed = {'cohort': COHORT['id'], 'version': version, 'source': names[primary], 'sources': [{'tag': t, 'name': names[t]} for t in tags],
            'start': min(dates), 'end': max(dates), 'courses': list(COHORT['courses']), 'modules': MODULES, 'lectures': lectures, 'classes': classes,
            'merge': {'primary': primary, 'rule': '칸별 과목·시간·강사: 한쪽만 있으면 그 값, 같으면 both, 다르면 첫 판 값 + review 표시',
                      'columns': {s['tag']: s['groups'] for s in sheets}, 'excluded': {s['tag']: s['excluded'] for s in sheets if s['excluded']},
                      'aligned': [{k: v for k, v in a.items() if k != 'moves'} for a in aligned],
                      'classes': {cid: {**{k: v for k, v in s.items() if k != 'conflicts'}, 'conflicts': dict(s['conflicts'])} for cid, s in summary.items()}}}
    reference = {'version': f'{version}-cohort{COHORT["id"]}', 'sources': [{'title': names[t]} for t in tags],
                 'courses': {course: [l['title'] for l in lectures if l['course'] == course and l['module'] == '직무특화' and l['title'] != UNDECIDED] for course in COHORT['courses']},
                 'facilities': facilities, 'patches': [], 'additions': []}

    if opts.get('out'):  # 시험·미리보기: 저장소 파일은 그대로 두고 다른 폴더에 씁니다(이력·비교 없음)
        out = pathlib.Path(opts['out']); out.mkdir(parents=True, exist_ok=True)
        (out / SEED.name).write_text(json.dumps(seed, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
        (out / REFERENCE.name).write_text(json.dumps(reference, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f'{out}에 {SEED.name}·{REFERENCE.name}를 썼습니다.')
        return
    if same_cohort:  # 이전 2기 시드와 비교(편집본 3자 병합과 같은 열쇠: 강의+일차, 휴일은 날짜)
        key = lambda e: f"h|{e['date']}|{e['title']}" if e['kind'] == 'holiday' else f"c|{e['lectureId']}|{e['day']}"
        fields, total = ('date', 'module', 'lectureId', 'title', 'day', 'hours', 'kind', 'online', 'instructorName', 'event'), 0
        for cid in CLASSES:
            before = {key(e): e for e in old['classes'].get(cid, {}).get('entries', [])}
            after = {key(e): e for e in classes[cid]['entries']}
            lines = [f"  + {after[k]['date']} {after[k]['title']}" for k in after.keys() - before.keys()]
            lines += [f"  - {before[k]['date']} {before[k]['title']}" for k in before.keys() - after.keys()]
            lines += [f"  ~ {after[k]['title']} {after[k]['day']}일차: " + ', '.join(f"{f} {before[k].get(f)}→{after[k].get(f)}" for f in fields if before[k].get(f) != after[k].get(f))
                      for k in after.keys() & before.keys() if any(before[k].get(f) != after[k].get(f) for f in fields)]
            if lines:
                total += len(lines); print(f'{cid}반 · {len(lines)}건'); print('\n'.join(sorted(lines)))
        if not total and not opts.get('force'):
            print('기존 시간표 기준과 달라진 수업이 없습니다(--force면 다시 씁니다).')
            return
    else:
        print(f"\n기존 시드({old.get('version', '없음')}, cohort {old.get('cohort', '1기')})를 {COHORT['label']} 시드로 바꿉니다.")
    if opts.get('check'):
        return
    if same_cohort and version == old['version'] and not opts.get('force'):
        sys.exit('버전이 기존과 같습니다. --version 으로 새 버전을 지정해 주세요.')
    HISTORY.mkdir(exist_ok=True)
    if same_cohort and not (HISTORY / f"{old['version']}.json").exists():  # 이전 판은 그 판에서 만든 편집본의 병합 기준
        (HISTORY / f"{old['version']}.json").write_text(SEED.read_text(encoding='utf-8'), encoding='utf-8')
    text = json.dumps(seed, ensure_ascii=False, separators=(',', ':'))
    SEED.write_text(text, encoding='utf-8'); (HISTORY / f'{version}.json').write_text(text, encoding='utf-8')
    REFERENCE.write_text(json.dumps(reference, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"{SEED.name}·timetable-seeds/{version}.json·{REFERENCE.name} 갱신 · {len(classes)}개 반 · {sum(len(c['entries']) for c in classes.values())}개 일정 · {seed['start']} ~ {seed['end']}")


if __name__ == '__main__':
    main()
